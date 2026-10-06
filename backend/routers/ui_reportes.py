"""Sección de Reportes: quién lleva qué (por persona) y qué actividades están cubiertas (por puesto),
con descarga a Excel, más la ficha de reestructuración en PDF para cada persona."""
import io
from datetime import date
from fastapi import APIRouter, Depends, Request, Response, HTTPException
from fastapi.templating import Jinja2Templates
from fastapi.responses import HTMLResponse
from sqlalchemy.orm import Session
from pathlib import Path
from database import get_db
from auth import verificar_acceso
from functions import calcular_cobertura, resumen_cobertura
from reparto import reporte_personas, actividades_a_cargo
import models as m

router = APIRouter(prefix="/ui/reportes", dependencies=[Depends(verificar_acceso)])
templates = Jinja2Templates(directory=str(Path(__file__).resolve().parents[1] / "templates"))

ESTADO_TXT = {"cubierta": "Cubierta", "cubierta_sin_etapa": "Cubierta, sin etapa", "parcial": "Parcial",
              "sin_asignar": "Sin asignar", "sin_actividades": "Sin actividades"}


# ───────────────────────── utilidades ─────────────────────────

def _xlsx(hojas, nombre):
    """hojas: lista de (titulo, encabezados, filas, anchos). Formato simple, sin adornos."""
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    wb = Workbook()
    wb.remove(wb.active)
    for titulo, encabezados, filas, anchos in hojas:
        ws = wb.create_sheet(titulo)
        ws.append(encabezados)
        for c in ws[1]:
            c.font = Font(bold=True, color="FFFFFF")
            c.fill = PatternFill("solid", fgColor="0A2640")
            c.alignment = Alignment(vertical="center", wrap_text=True)
        for f in filas:
            ws.append(f)
        for i, ancho in enumerate(anchos, start=1):
            ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = ancho
        for fila in ws.iter_rows(min_row=2):
            for c in fila:
                c.alignment = Alignment(vertical="top", wrap_text=True)
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
    buf = io.BytesIO()
    wb.save(buf)
    return Response(content=buf.getvalue(), media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    headers={"Content-Disposition": f'attachment; filename="{nombre}"'})


def _planta_nombre(db, planta_id):
    return db.query(m.Planta).get(planta_id).nombre if planta_id else None


# ───────────────────────── índice ─────────────────────────

@router.get("", response_class=HTMLResponse)
def indice(request: Request):
    return templates.TemplateResponse(request, "reportes.html", {})


# ───────────────────────── asignación por persona ─────────────────────────

@router.get("/personas", response_class=HTMLResponse)
def reporte_por_persona(request: Request, planta_id: int = None, incluir_bajas: bool = False, db: Session = Depends(get_db)):
    filas = reporte_personas(db, planta_id, incluir_bajas)
    con_asig = [f for f in filas if not f["sin_asignacion"]]
    return templates.TemplateResponse(request, "reporte_personas.html", {
        "filas": filas, "planta_id": planta_id, "incluir_bajas": incluir_bajas,
        "plantas": db.query(m.Planta).order_by(m.Planta.id).all(),
        "n_personas": len(filas), "n_con_asignacion": len(con_asig),
        "n_sin_asignacion": len(filas) - len(con_asig),
        "total_a_cargo": sum(f["n_total"] for f in filas), "total_cedidas": sum(f["n_cedidas"] for f in filas),
    })


@router.get("/personas.xlsx")
def reporte_por_persona_xlsx(planta_id: int = None, incluir_bajas: bool = False, db: Session = Depends(get_db)):
    filas = reporte_personas(db, planta_id, incluir_bajas)
    resumen, detalle = [], []
    for f in filas:
        resumen.append([f["persona"], ", ".join(f["puestos"]) or "—", ", ".join(f["plantas"]) or "—",
                        f["n_propias"], f["n_recibidas"], f["n_cedidas"], f["n_total"], f["n_criticas"]])
        for b in f["bloques_propios"]:
            for a in b["propias"]:
                detalle.append([f["persona"], "Propia (de su puesto)", f"{b['puesto_codigo']} — {b['puesto_nombre']}", a["codigo"],
                                a["descripcion"], "Sí" if a["es_critica"] else "", a["frecuencia"] or "", b["planta"], ""])
            for c in b["cedidas"]:
                a = c["actividad"]
                detalle.append([f["persona"], "Cedida (la lleva otra persona)", f"{b['puesto_codigo']} — {b['puesto_nombre']}", a["codigo"],
                                a["descripcion"], "Sí" if a["es_critica"] else "", a["frecuencia"] or "", b["planta"],
                                "La lleva: " + (", ".join(c["a_cargo_de"]) if c["a_cargo_de"] else "nadie (sin cubrir)")])
        for b in f["bloques_recibidos"]:
            for a in b["actividades"]:
                detalle.append([f["persona"], "Recibida (de otro puesto)", f"{b['puesto_codigo'] or '—'} — {b['puesto_nombre']}", a["codigo"],
                                a["descripcion"], "Sí" if a["es_critica"] else "", a["frecuencia"] or "", b["planta"],
                                ("Titular del puesto: " + ", ".join(b["titulares"])) if b["titulares"] else ""])
    return _xlsx([
        ("Resumen por persona", ["Persona", "Puesto(s)", "Planta(s)", "Propias", "Recibidas", "Cedidas", "Total a cargo", "Críticas"],
         resumen, [30, 22, 30, 10, 11, 10, 14, 10]),
        ("Detalle", ["Persona", "Tipo", "Puesto de origen", "Código", "Actividad", "Crítica", "Frecuencia", "Planta", "Observación"],
         detalle, [28, 28, 34, 11, 70, 9, 14, 24, 36]),
    ], f"Asignacion_por_persona_{date.today().isoformat()}.xlsx")


# ───────────────────────── ficha de reestructuración (PDF) ─────────────────────────

def _ctx_ficha_pdf(db: Session, persona_id: int):
    persona = db.query(m.Persona).get(persona_id)
    if not persona:
        raise HTTPException(404, "Persona no encontrada")
    a_cargo = actividades_a_cargo(db, persona_id)

    def etiqueta(cod):
        p = db.query(m.Puesto).filter(m.Puesto.codigo == cod).first()
        return f"{cod} — {p.nombre}" if p else cod

    def referencias(cod):
        mecs = [f"{mp.mecanismo.codigo} {mp.mecanismo.nombre} ({mp.rol})"
                for mp in db.query(m.MecanismoPuesto).filter(m.MecanismoPuesto.puesto_codigo == cod)]
        procs = [f"{pr.codigo} {pr.nombre}" for pr in db.query(m.Procedimiento).join(
                    m.ProcedimientoPuesto, m.ProcedimientoPuesto.procedimiento_id == m.Procedimiento.id)
                 .filter(m.ProcedimientoPuesto.puesto_codigo == cod, m.Procedimiento.vigente == 1)]
        inds = [f"{i.codigo} {i.nombre}" for i in db.query(m.Indicador).filter(
                    m.Indicador.activo == 1,
                    (m.Indicador.responsable_carga_puesto_codigo == cod) | (m.Indicador.responsable_uso_puesto_codigo == cod))]
        return {"mecanismos": mecs, "procedimientos": procs, "indicadores": inds}

    puestos = []
    for b in a_cargo["bloques_propios"]:
        p = db.query(m.Puesto).filter(m.Puesto.codigo == b["puesto_codigo"]).first()
        rel = db.query(m.RelacionReporte).filter(m.RelacionReporte.puesto_codigo == b["puesto_codigo"]).all()
        puestos.append({"codigo": b["puesto_codigo"], "nombre": b["puesto_nombre"], "planta": b["planta"], "etapa": b["etapa"],
                        "proposito": p.proposito if p else None,
                        "reporta_a": [etiqueta(r.relacionado_codigo) for r in rel if r.tipo == "jerarquico"],
                        "coordina_con": [etiqueta(r.relacionado_codigo) for r in rel if r.tipo == "funcional"],
                        **referencias(b["puesto_codigo"])})
    origenes = []
    for b in a_cargo["bloques_recibidos"]:
        if b["puesto_codigo"] and b["puesto_codigo"] not in {o["codigo"] for o in origenes}:
            origenes.append({"codigo": b["puesto_codigo"], "nombre": b["puesto_nombre"], **referencias(b["puesto_codigo"])})
    return {"persona": persona, "a_cargo": a_cargo, "puestos": puestos, "origenes": origenes,
            "fecha": date.today().strftime("%d/%m/%Y")}


@router.get("/personas/{persona_id}.pdf")
def ficha_persona_pdf(persona_id: int, db: Session = Depends(get_db)):
    from xhtml2pdf import pisa
    ctx = _ctx_ficha_pdf(db, persona_id)
    html = templates.get_template("pdf_ficha_persona.html").render(**ctx)
    buf = io.BytesIO()
    pisa.CreatePDF(html, dest=buf)
    nombre = f"Ficha_{ctx['persona'].apellido}_{ctx['persona'].nombre}".replace(" ", "_").replace("(", "").replace(")", "")
    return Response(content=buf.getvalue(), media_type="application/pdf",
                    headers={"Content-Disposition": f'inline; filename="{nombre}.pdf"'})


# ───────────────────────── cobertura por puesto ─────────────────────────

def _filas_cobertura(db, planta_id, solo_pendientes):
    filas = calcular_cobertura(db, planta_id=planta_id)
    if solo_pendientes:
        filas = [f for f in filas if f["estado"] in ("sin_asignar", "parcial")]
    return filas


@router.get("/cobertura", response_class=HTMLResponse)
def reporte_cobertura(request: Request, planta_id: int = None, solo_pendientes: bool = False, db: Session = Depends(get_db)):
    todas = calcular_cobertura(db, planta_id=planta_id)
    filas = [f for f in todas if f["estado"] in ("sin_asignar", "parcial")] if solo_pendientes else todas
    return templates.TemplateResponse(request, "reporte_cobertura.html", {
        "filas": filas, "resumen": resumen_cobertura(todas), "planta_id": planta_id, "solo_pendientes": solo_pendientes,
        "plantas": db.query(m.Planta).order_by(m.Planta.id).all(), "ESTADO_TXT": ESTADO_TXT,
        "n_sin_cubrir": sum(f["n_sin_cubrir"] for f in todas if f["estado"] != "sin_asignar"),
        "n_cedidas": sum(f["n_cedidas"] for f in todas),
    })


@router.get("/cobertura.xlsx")
def reporte_cobertura_xlsx(planta_id: int = None, solo_pendientes: bool = False, db: Session = Depends(get_db)):
    filas = _filas_cobertura(db, planta_id, solo_pendientes)
    cob, sin = [], []
    for f in filas:
        etapa = f["etapa"] if f["etapa"] and f["etapa"] != "sin_etapa" else ""
        cob.append([f["puesto_codigo"], f["puesto_nombre"], f["planta_nombre"], etapa, ESTADO_TXT.get(f["estado"], f["estado"]),
                    f["n_actividades"], f["n_por_titular"], f["n_cedidas"], f["n_por_suelta"], f["n_sin_cubrir"],
                    ", ".join(r["persona"] + (" (perfil)" if r["via"] == "perfil" else "") for r in f["responsables"]) or "—"])
        for a in f["actividades_sin_asignar"]:
            sin.append([f["puesto_codigo"], f["puesto_nombre"], f["planta_nombre"], etapa, a["codigo"], a["descripcion"], "Sí" if a["es_critica"] else ""])
    return _xlsx([
        ("Cobertura por puesto", ["Puesto", "Nombre", "Planta", "Etapa", "Estado", "Actividades", "Las lleva el titular",
                                  "Cedidas", "Las lleva otra persona", "Sin cubrir", "Responsables"],
         cob, [10, 36, 26, 11, 18, 12, 14, 10, 14, 11, 44]),
        ("Actividades sin cubrir", ["Puesto", "Nombre", "Planta", "Etapa", "Código", "Actividad", "Crítica"],
         sin, [10, 32, 26, 11, 11, 80, 9]),
    ], f"Cobertura_por_puesto_{date.today().isoformat()}.xlsx")
